from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree
import zipfile

from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


MAX_EXTRACTED_CHARS = 2_000_000


class UnsupportedDocumentError(Exception):
    pass


@dataclass
class ExtractionResult:
    text: str
    extractor: str
    truncated: bool


def extract_document(path: Path, extension: str) -> ExtractionResult:
    if extension == ".pdf":
        text = _extract_pdf(path)
        extractor = "pypdf"
    elif extension == ".hwpx":
        text = _extract_hwpx(path)
        extractor = "hwpx-xml"
    elif extension == ".pptx":
        text = _extract_pptx(path)
        extractor = "python-pptx"
    elif extension == ".xlsx":
        text = _extract_xlsx(path)
        extractor = "openpyxl"
    elif extension in {".hwp", ".ppt", ".xls"}:
        raise UnsupportedDocumentError(
            "HWP/PPT/XLS 구형 바이너리 형식은 전용 파서 연결 후 지원할 수 있습니다."
        )
    else:
        raise UnsupportedDocumentError("지원하지 않는 문서 형식입니다.")

    normalized = _normalize_text(text)
    truncated = len(normalized) > MAX_EXTRACTED_CHARS
    if truncated:
        normalized = normalized[:MAX_EXTRACTED_CHARS]
    return ExtractionResult(
        text=normalized,
        extractor=extractor,
        truncated=truncated,
    )


def _extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_hwpx(path: Path) -> str:
    sections: list[str] = []
    with zipfile.ZipFile(path) as archive:
        names = sorted(
            name
            for name in archive.namelist()
            if name.startswith("Contents/section") and name.endswith(".xml")
        )
        for name in names:
            root = ElementTree.fromstring(archive.read(name))
            sections.append(" ".join(root.itertext()))
    return "\n".join(sections)


def _extract_pptx(path: Path) -> str:
    presentation = Presentation(str(path))
    slides: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        paragraphs: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                paragraphs.append(shape.text)
        slides.append(f"[슬라이드 {index}]\n" + "\n".join(paragraphs))
    return "\n".join(slides)


def _extract_xlsx(path: Path) -> str:
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        sheets: list[str] = []
        for worksheet in workbook.worksheets:
            rows: list[str] = [f"[시트: {worksheet.title}]"]
            for row in worksheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None]
                if values:
                    rows.append("\t".join(values))
            sheets.append("\n".join(rows))
        return "\n".join(sheets)
    finally:
        workbook.close()


def _normalize_text(text: str) -> str:
    return "\n".join(line.strip() for line in text.splitlines() if line.strip())
