import base64
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document as WordDocument
from pptx import Presentation

from app.extraction import UnsupportedDocumentError, extract_document, odf_content_to_text

# 1x1 PNG 픽셀. 실제 이미지 내용은 중요하지 않다 - OCR 엔진 자체는 아래에서
# mock으로 대체하고, "이미지가 있으면 그 bytes가 OCR로 넘어가 결과가 본문에
# 합쳐지는지"만 검증한다.
_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


class TextExtractionTests(unittest.TestCase):
    def test_extracts_markdown_as_plain_text(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "policy.md"
            path.write_text("# 보안 정책\n\n외부 AI 전송 전 검사", encoding="utf-8")

            result = extract_document(path, ".md")

            self.assertEqual(result.extractor, "plain-text")
            self.assertIn("보안 정책", result.text)
            self.assertFalse(result.truncated)

    def test_extracts_cp949_text_file(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "memo.txt"
            path.write_bytes("기관 내부 문서".encode("cp949"))

            result = extract_document(path, ".txt")

            self.assertEqual(result.text, "기관 내부 문서")

    def test_extracts_docx_paragraphs_and_tables(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.docx"
            document = WordDocument()
            document.add_paragraph("외부 전송 전 보안 검사가 필요합니다.")
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "구분"
            table.rows[0].cells[1].text = "기밀"
            document.save(str(path))

            result = extract_document(path, ".docx")

            self.assertEqual(result.extractor, "python-docx")
            self.assertIn("외부 전송 전 보안 검사가 필요합니다.", result.text)
            self.assertIn("구분\t기밀", result.text)

    def test_extracts_text_from_image_pasted_into_pptx(self):
        """QA가 보고한 블라인드 스팟: 기밀 내용을 캡처해 이미지로 PPT에 끼워
        넣으면, 텍스트 상자만 읽는 python-pptx로는 전혀 보이지 않아 보안
        검사를 그대로 통과했다. 이미지 OCR을 붙인 뒤 다시 검증한다."""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "capture.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_picture(io.BytesIO(_TINY_PNG), 0, 0)
            presentation.save(str(path))

            fake_reader = type("FakeReader", (), {"readtext": staticmethod(lambda *a, **k: ["기밀 문서 캡처본"])})()
            with patch("app.extraction._get_ocr_reader", return_value=fake_reader):
                result = extract_document(path, ".pptx")

            self.assertEqual(result.extractor, "python-pptx+easyocr")
            self.assertIn("기밀 문서 캡처본", result.text)

    def test_pptx_without_images_does_not_report_ocr_usage(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "plain.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            textbox = slide.shapes.add_textbox(0, 0, 1000, 1000)
            textbox.text_frame.text = "일반 텍스트 슬라이드"
            presentation.save(str(path))

            result = extract_document(path, ".pptx")

            self.assertEqual(result.extractor, "python-pptx")
            self.assertIn("일반 텍스트 슬라이드", result.text)

    def test_extracts_text_from_image_pasted_into_docx(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "capture.docx"
            document = WordDocument()
            document.add_paragraph("첨부 캡처 화면:")
            document.add_picture(io.BytesIO(_TINY_PNG))
            document.save(str(path))

            fake_reader = type("FakeReader", (), {"readtext": staticmethod(lambda *a, **k: ["주민등록번호 캡처됨"])})()
            with patch("app.extraction._get_ocr_reader", return_value=fake_reader):
                result = extract_document(path, ".docx")

            self.assertEqual(result.extractor, "python-docx+easyocr")
            self.assertIn("주민등록번호 캡처됨", result.text)

    def test_hwp_with_invalid_container_is_unsupported(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "broken.hwp"
            path.write_bytes(b"not a real hwp file" * 10)

            with self.assertRaises(UnsupportedDocumentError):
                extract_document(path, ".hwp")

    def test_extracts_csv_as_plain_text(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "rows.csv"
            path.write_text("이름,등급\n기획서,S\n", encoding="utf-8")

            result = extract_document(path, ".csv")

            self.assertEqual(result.extractor, "plain-text")
            self.assertIn("기획서,S", result.text)

    def test_extracts_html_and_skips_script_and_style(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "page.html"
            path.write_text(
                "<html><head><style>body{color:red}</style></head>"
                "<body><h1>공지</h1><p>외부 전송 전 검사가 필요합니다.</p>"
                "<script>alert('should not appear')</script></body></html>",
                encoding="utf-8",
            )

            result = extract_document(path, ".html")

            self.assertEqual(result.extractor, "html-parser")
            self.assertIn("공지", result.text)
            self.assertIn("외부 전송 전 검사가 필요합니다.", result.text)
            self.assertNotIn("alert", result.text)
            self.assertNotIn("color:red", result.text)


class OdfContentTextTests(unittest.TestCase):
    def test_paragraphs_and_table_cells_are_flattened(self):
        content_xml = """<?xml version="1.0" encoding="UTF-8"?>
        <office:document-content
            xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
            xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
            xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0">
          <office:body>
            <office:text>
              <text:p>외부 전송 전 검사가 필요합니다.</text:p>
              <table:table>
                <table:table-row>
                  <table:table-cell><text:p>구분</text:p></table:table-cell>
                  <table:table-cell><text:p>기밀</text:p></table:table-cell>
                </table:table-row>
              </table:table>
            </office:text>
          </office:body>
        </office:document-content>
        """.encode("utf-8")

        text = odf_content_to_text(content_xml)

        self.assertIn("외부 전송 전 검사가 필요합니다.", text)
        self.assertIn("구분", text)
        self.assertIn("기밀", text)


if __name__ == "__main__":
    unittest.main()
